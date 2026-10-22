"""Offline batch checks: python probe/check_batch.py. No real mail or WP calls."""

from email.message import EmailMessage
from io import BytesIO
import os
from pathlib import Path
import runpy
import unittest
from unittest.mock import MagicMock, patch

import check_logging as harness
from docx import Document
import eml_links
import image_manager
import integrity_check
import log
import mail_downloader as mail
import main
import news_parser
import prelaunch


def document_bytes(title="Заголовок", text="Содержимое новости."):
    document = Document()
    document.add_paragraph(title)
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def email_bytes(attachments=(), text="Материалы для публикации."):
    message = EmailMessage()
    message["Subject"] = "ЗГП"
    message["From"] = "sender@example.invalid"
    message["To"] = "news@example.invalid"
    message.set_content(text)
    for name, data in attachments:
        message.add_attachment(data, maintype="application", subtype="octet-stream", filename=name)
    return message.as_bytes()


class BatchChecks(unittest.TestCase):
    def setUp(self):
        harness.LoggingChecks.setUp(self)
        self.box = {}
        self.seen = set()
        self.copied = []
        self.events = []
        self.refuse_unseen = None
        self.on_fetch = None
        self.stack.enter_context(patch.object(main, "OUTPUT_DIR", str(self.root / "extracted_images")))
        self.stack.enter_context(patch.object(mail, "IMAP_OUTPUT_DIR", str(self.root / "eml")))
        old_cwd = Path.cwd()
        os.chdir(self.root)
        self.addCleanup(os.chdir, old_cwd)
        self.connection.uid.side_effect = self.uid
        self.search = self.stack.enter_context(patch.object(mail, "search_unseen_news",
            side_effect=lambda _: [(uid, "ЗГП") for uid in sorted(self.box, key=int) if uid not in self.seen]))
        # Exercise real MIME/DOCX parsing and image matching; only Pandoc is substituted.
        self.stack.enter_context(patch.object(news_parser, "convert_docx_to_html",
            return_value="<p>Заголовок</p><p>Содержимое новости.</p>"))
        self.posts = self.stack.enter_context(patch.object(main, "create_post",
            side_effect=lambda n: {"id": 100, "link": "https://example.invalid/news/100", "status": "draft"}))
        self.upload = self.stack.enter_context(patch.object(main, "upload_media", return_value={"wp_id": 42}))
        self.preview_upload = self.stack.enter_context(patch.object(main, "upload_media_bytes", return_value={"wp_id": 43}))
        self.default = self.stack.enter_context(patch.object(main, "get_default_media_id", return_value=77))

    def uid(self, command, uid, *args):
        if command == "fetch":
            if args[0] == "(FLAGS)":
                flags = b"\\Seen" if uid in self.seen else b""
                return "OK", [b"1 (UID " + uid.encode() + b" FLAGS (" + flags + b"))"]
            self.assertEqual(args[0], "(BODY.PEEK[])")
            if self.on_fetch is not None:
                self.on_fetch(uid)
            return "OK", [(b"1 (BODY[])", self.box[uid])]
        if command == "store":
            if args[1] == "(\\Seen)":
                if args[0] == "-FLAGS":
                    self.events.append(("unseen", uid))
                    if uid != self.refuse_unseen:
                        self.seen.discard(uid)
                else:
                    self.seen.add(uid)
            return "OK", []
        if command == "copy":
            self.copied.append(uid)
            return "OK", []
        raise AssertionError((command, uid, args))

    def add(self, uid, attachments=(), text="Материалы."):
        self.box[str(uid)] = email_bytes(attachments, text)

    def assert_clean(self):
        for folder in ("eml", "extracted_images"):
            directory = self.root / folder
            self.assertFalse(directory.exists() and any(directory.iterdir()))
        self.assertFalse(log._logger.handlers)
        lock = prelaunch.acquire_process_lock()
        self.assertIsNotNone(lock)
        prelaunch.release_process_lock(lock)

    def test_docx_and_photo_in_different_emails_share_one_batch(self):
        self.add(1, [("Уборка.docx", document_bytes())])
        self.add(2, [("Уборка_1.jpg", b"photo")])
        prelaunch.main()
        self.posts.assert_called_once()
        self.upload.assert_called_once()
        self.default.assert_not_called()
        self.assertEqual(self.copied, ["1", "2"])
        self.search.assert_called_once()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_deferred_batch_is_redownloaded_with_later_email(self):
        self.add(1, [("Уборка.docx", document_bytes())])
        original_open = Path.open

        def observe_flag(path, *args, **kwargs):
            if path == prelaunch.DELAYED_PUBLISH_FILE and args and args[0] == "x":
                self.assertEqual(self.seen, set())
                self.assertEqual(self.events, [("unseen", "1")])
            return original_open(path, *args, **kwargs)

        with patch.object(Path, "open", autospec=True, side_effect=observe_flag):
            prelaunch.main()
        self.assertTrue(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.posts.assert_not_called()
        self.default.assert_not_called()
        self.assert_clean()
        self.add(2, [("Уборка_1.jpg", b"late photo")])
        prelaunch.main()
        self.posts.assert_called_once()
        self.assertEqual(self.copied, ["1", "2"])
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_repeat_uses_existing_default_once_for_two_news(self):
        self.add(1, [("Уборка.docx", document_bytes("Первая")),
                     ("Спорт.docx", document_bytes("Вторая"))])
        prelaunch.main()
        self.assertTrue(prelaunch.DELAYED_PUBLISH_FILE.exists())
        prelaunch.main()
        self.assertEqual(self.posts.call_count, 2)
        self.default.assert_called_once()
        self.upload.assert_not_called()
        for call in self.posts.call_args_list:
            self.assertEqual(call.args[0]["featured_media_id"], 77)
            self.assertEqual(call.args[0]["gallery_media_ids"], [])
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_duplicate_names_keep_first_bytes_without_postponement(self):
        self.add(1, [("Уборка.docx", document_bytes("Первая версия")), ("Уборка.jpg", b"first")])
        self.add(2, [("УБОРКА.DOCX", document_bytes("Вторая версия")), ("УБОРКА.JPG", b"different")])
        prelaunch.main()
        self.posts.assert_called_once()
        self.assertEqual(self.posts.call_args.args[0]["title"], "Первая версия")
        self.upload.assert_called_once()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        text = next(self.logs.glob("*.txt")).read_text(encoding="utf-8")
        self.assertEqual(text.count("Дубль вложения пропущен"), 2)
        self.assert_clean()

    def test_gallery_reuses_featured_upload(self):
        self.add(1, [("Уборка.docx", document_bytes()), ("Уборка_1.jpg", b"first"),
                     ("Уборка_2.jpg", b"second")])
        prelaunch.main()
        self.assertEqual(self.upload.call_count, 2)
        news = self.posts.call_args.args[0]
        self.assertEqual(len(news["gallery_media_ids"]), 2)
        self.assertEqual(news["featured_media_id"], news["gallery_media_ids"][0])
        self.assert_clean()

    def test_default_lookup_failure_skips_only_affected_news(self):
        self.add(1, [("Спорт.docx", document_bytes()), ("Уборка.docx", document_bytes()),
                     ("Уборка.jpg", b"photo")])
        prelaunch.DELAYED_PUBLISH_FILE.touch()
        self.default.side_effect = RuntimeError("default image not found")
        prelaunch.main()
        self.posts.assert_called_once()
        self.assertEqual(self.posts.call_args.args[0]["document_filename"], "Уборка.docx")
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assertEqual(self.copied, [])
        self.assert_clean()

    def test_orphan_photo_defers_once_then_publishes_available_news(self):
        self.add(1, [("Уборка.docx", document_bytes()), ("Уборка.jpg", b"photo"), ("Спорт.jpg", b"orphan")])
        prelaunch.main()
        self.posts.assert_not_called()
        self.assertTrue(prelaunch.DELAYED_PUBLISH_FILE.exists())
        prelaunch.main()
        self.posts.assert_called_once()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_vk_without_photo_does_not_defer_and_skips_if_preview_missing(self):
        self.add(1, [("Видео.docx", document_bytes(text="https://vk.ru/wall-1_2")),
                     ("Уборка.docx", document_bytes()), ("Уборка.jpg", b"photo")])
        with patch.object(main, "get_vk_video_from_post", return_value={"player": "https://example.invalid/player"}):
            prelaunch.main()
        self.posts.assert_called_once()
        self.assertEqual(self.posts.call_args.args[0]["document_filename"], "Уборка.docx")
        self.default.assert_not_called()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assertEqual(self.copied, [])  # Partial batch stays SEEN, avoiding automatic reposts.
        self.assertEqual(self.seen, {"1"})
        self.assert_clean()

    def test_vk_preview_replaces_default_and_docx_link_has_priority(self):
        self.add(1, [("Видео.docx", document_bytes(text="https://vk.ru/wall-1_2"))],
                 text="https://vk.ru/wall-3_4\nВидео.docx")
        with patch.object(main, "get_vk_video_from_post", return_value={"preview": "https://example.invalid/preview"}) as vk, \
             patch.object(main, "download_vk_preview", return_value=(b"preview", "image/jpeg", ".jpg")):
            prelaunch.main()
        vk.assert_called_once_with("https://vk.ru/wall-1_2")
        self.posts.assert_called_once()
        self.assertEqual(self.posts.call_args.args[0]["featured_media_id"], 43)
        self.default.assert_not_called()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_vk_reference_in_later_email_can_match_previous_docx(self):
        self.add(1, [("Видео.docx", document_bytes())])
        self.add(2, text="https://vk.ru/wall-3_4\nВидео.docx")
        with patch.object(main, "get_vk_video_from_post", return_value={"preview": "https://example.invalid/preview"}), \
             patch.object(main, "download_vk_preview", return_value=(b"preview", "image/jpeg", ".jpg")):
            prelaunch.main()
        self.posts.assert_called_once()
        self.default.assert_not_called()
        self.assert_clean()

    def test_email_text_boundaries_are_preserved(self):
        self.add(1, [("Видео.docx", document_bytes())], text="https://vk.ru/wall-3_4")
        self.add(2, text="Видео.docx")
        with self.assertRaisesRegex(RuntimeError, "Не удалось сопоставить"):
            prelaunch.main()
        self.posts.assert_not_called()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_later_arrival_does_not_extend_snapshot(self):
        self.add(1, [("Уборка.docx", document_bytes()), ("Уборка.jpg", b"photo")])
        self.on_fetch = lambda _: self.add(2, [("Спорт.docx", document_bytes()), ("Спорт.jpg", b"photo")])
        prelaunch.main()
        self.assertEqual(self.posts.call_count, 1)
        self.assertEqual(self.seen, {"1"})
        self.assert_clean()

    def test_terminal_download_error_stops_whole_batch_and_cleans(self):
        self.add(1, [("Уборка.docx", document_bytes())])
        self.add(2, [("Уборка.jpg", b"photo")])
        self.add(3, [("Спорт.docx", document_bytes())])
        original = mail.download_by_uid
        calls = []

        def download(connection, uid):
            calls.append(uid)
            if uid == "2":
                raise OSError("download failed")
            return original(connection, uid)

        with patch.object(mail, "download_by_uid", side_effect=download), patch.object(mail.time, "sleep"):
            with self.assertRaisesRegex(RuntimeError, "не удалось скачать UID 2"):
                prelaunch.main()
        self.assertEqual(calls, ["1"] + ["2"] * mail.IMAP_RETRY_COUNT)
        self.posts.assert_not_called()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_corrupt_eml_stops_preparation_without_delay(self):
        self.box["1"] = b"Subject: ZGP\nContent-Type: multipart/mixed; boundary=missing\n\ninvalid"
        with self.assertRaises(Exception):
            prelaunch.main()
        self.posts.assert_not_called()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assertEqual(self.seen, {"1"})
        self.assert_clean()

    def test_corrupt_base64_attachment_stops_preparation(self):
        self.box["1"] = (
            b'Subject: ZGP\r\nContent-Type: multipart/mixed; boundary="part"\r\n\r\n'
            b'--part\r\nContent-Type: application/octet-stream\r\n'
            b'Content-Disposition: attachment; filename="news.docx"\r\n'
            b'Content-Transfer-Encoding: base64\r\n\r\n!bad!\r\n--part--\r\n'
        )
        with self.assertRaises(Exception):
            prelaunch.main()
        self.posts.assert_not_called()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_unconfirmed_unseen_does_not_create_flag_or_publish(self):
        self.add(1, [("Уборка.docx", document_bytes())])
        self.refuse_unseen = "1"
        with self.assertRaisesRegex(RuntimeError, "не подтвердил UNSEEN"):
            prelaunch.main()
        self.posts.assert_not_called()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_flag_creation_failure_has_no_mail_rollback(self):
        self.add(1, [("Уборка.docx", document_bytes())])
        original = Path.open

        def fail_flag(path, *args, **kwargs):
            if path == prelaunch.DELAYED_PUBLISH_FILE and args and args[0] == "x":
                raise OSError("flag creation failed")
            return original(path, *args, **kwargs)

        with patch.object(Path, "open", autospec=True, side_effect=fail_flag):
            with self.assertRaisesRegex(OSError, "flag creation failed"):
                prelaunch.main()
        self.assertEqual(self.seen, set())
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.posts.assert_not_called()
        self.assert_clean()

    def test_repeat_flag_is_consumed_on_error_and_empty_mailbox(self):
        prelaunch.DELAYED_PUBLISH_FILE.touch()
        with patch.object(mail, "search_unseen_news", side_effect=OSError("IMAP down")):
            with self.assertRaisesRegex(OSError, "IMAP down"):
                prelaunch.main()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        prelaunch.DELAYED_PUBLISH_FILE.touch()
        prelaunch.main()
        self.assertFalse(prelaunch.DELAYED_PUBLISH_FILE.exists())
        self.assert_clean()

    def test_lock_is_held_through_pipeline_and_stale_file_is_harmless(self):
        Path(prelaunch.LOCK_FILE).touch()

        def run(**kwargs):
            self.assertIsNone(prelaunch.acquire_process_lock())
            return 0, 0, 0, 0

        with patch.object(main, "run_pipeline", side_effect=run):
            prelaunch.main()
        self.assert_clean()

    def test_legacy_script_entry_uses_prelaunch(self):
        with patch.object(prelaunch, "main") as launch:
            runpy.run_path(str(harness.ROOT / "main.py"), run_name="__main__")
        launch.assert_called_once_with()

    def test_integrity_combinations(self):
        for with_image in (False, True):
            for with_vk in (False, True):
                with self.subTest(image=with_image, vk=with_vk):
                    report = integrity_check.check_integrity(
                        [{"filename": "Уборка.docx"}],
                        [{"filename": "Уборка.jpg"}] if with_image else [],
                        ["Уборка.docx"] if with_vk else [],
                    )
                    self.assertEqual(report["incomplete"], not with_image and not with_vk)

    def test_missing_flags_is_not_accepted_as_unseen(self):
        with patch.object(self.connection, "uid", return_value=("OK", [None])):
            with self.assertRaisesRegex(RuntimeError, "не вернул FLAGS"):
                mail.mark_email_unseen("1")

    def test_media_lookup_paginates_and_requires_exact_filename(self):
        first = MagicMock()
        first.json.return_value = [{"id": 1, "source_url": "https://example.invalid/s600xu_2x-1.webp"}]
        first.headers = {"X-WP-TotalPages": "2"}
        second = MagicMock()
        second.json.return_value = [{"id": 72, "source_url": "https://example.invalid/s600xu_2x.webp?x=1"}]
        second.headers = {"X-WP-TotalPages": "2"}
        with patch.object(image_manager.requests, "get", side_effect=[first, second]) as get:
            self.assertEqual(image_manager.get_default_media_id(), 72)
        self.assertEqual([c.kwargs["params"]["page"] for c in get.call_args_list], [1, 2])

    def test_media_lookup_raises_when_exact_file_is_absent(self):
        response = MagicMock()
        response.json.return_value = [{"id": 1, "source_url": "https://example.invalid/s600xu_2x-1.webp"}]
        response.headers = {"X-WP-TotalPages": "1"}
        with patch.object(image_manager.requests, "get", return_value=response):
            with self.assertRaisesRegex(RuntimeError, "не найдена"):
                image_manager.get_default_media_id()


if __name__ == "__main__":
    unittest.main(verbosity=2)
